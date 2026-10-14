"""python/export_xlsx.py on a temp project built from the fictional scoring fixtures."""
import copy
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / 'python' / 'export_xlsx.py'
FIX = ROOT / 'test' / 'fixtures' / 'scoring'


def make_project(tmp_path, mutate=None):
    (tmp_path / 'data').mkdir()
    shutil.copy(FIX / 'project.json', tmp_path / 'project.json')
    doc = json.loads((FIX / 'scored.json').read_text())
    if mutate:
        mutate(doc['items'])
    (tmp_path / 'data' / 'scored.json').write_text(json.dumps(doc))
    return tmp_path, doc['items']


def export(project):
    return subprocess.run([sys.executable, str(SCRIPT), str(project)], capture_output=True, text=True)


def headers(ws):
    return [c.value for c in ws[1]]


def col(ws, label):
    return headers(ws).index(label) + 1


@pytest.fixture
def built(tmp_path):
    project, rows = make_project(tmp_path)
    r = export(project)
    assert r.returncode == 0, r.stderr
    return project, rows, load_workbook(project / 'out' / 'leads.xlsx')


def test_sheets_and_row_counts(built):
    _p, rows, wb = built
    assert wb.sheetnames == ['Call list', 'All leads', 'Disqualified', 'Read me']
    live = [r for r in rows if not r['disqualified']]
    call = [r for r in live if r['tier'] in ('A', 'B')]
    assert wb['Call list'].max_row - 1 == len(call) >= 1
    assert wb['All leads'].max_row - 1 == len(live)
    assert wb['Disqualified'].max_row - 1 == len(rows) - len(live) == 5


def test_headers_unique_and_complete(built):
    _p, _rows, wb = built
    call, full = headers(wb['Call list']), headers(wb['All leads'])
    for need in ['Tier', 'Priority', 'Need', 'Interest', 'Business', 'Category', 'City', 'Sources', 'Owner', 'Phone', 'Phone check',
                 'Email', 'Website', 'SEO score', 'SEO band', 'SEO top issues', 'Instagram', 'Google rating', 'Google reviews',
                 'Runs ads', 'Incumbent agency', 'Defects', 'Why they need you', 'Why they will care', 'Risk', 'Channel',
                 'Outreach rule', 'Called on', 'Outcome', 'Next step', 'Next step date', 'Objection', 'Notes']:
        assert need in call, need
    assert len(full) == len(set(full)), 'duplicated column'
    assert full[:len(call)] == call
    assert 'Why it was cut' in headers(wb['Disqualified'])


def test_styling_links_and_values(built):
    project, rows, wb = built
    ws = wb['Call list']
    assert ws.freeze_panes == 'F2'
    assert ws.auto_filter.ref.startswith('A1:')
    assert ws.cell(1, 1).fill.fgColor.rgb.endswith('2E5E4E')            # agency.color
    top = rows[0]
    assert ws.cell(2, col(ws, 'Business')).value == top['name']
    assert ws.cell(2, col(ws, 'Tier')).fill.fgColor.rgb.endswith('C8E6C9')
    site = ws.cell(2, col(ws, 'Website'))
    assert site.hyperlink and site.hyperlink.target == top['website']
    assert ws.cell(2, col(ws, 'Instagram')).hyperlink is not None
    assert ws.cell(2, col(ws, 'Google Maps')).hyperlink is not None
    assert ws.cell(2, col(ws, 'Phone check')).value.startswith('MISMATCH')
    assert ws.column_dimensions['E'].width >= 30
    assert ws.cell(2, col(ws, 'Why they need you')).alignment.wrap_text


def test_not_checked_and_rtl(built):
    _p, _rows, wb = built
    ws = wb['All leads']
    by_name = {ws.cell(r, col(ws, 'Business')).value: r for r in range(2, ws.max_row + 1)}
    lark = by_name['Lark Salon']
    assert ws.cell(lark, col(ws, 'Google rating')).value == 'not checked'
    assert ws.cell(lark, col(ws, 'SEO score')).value == 'not checked'
    rtl = next(r for n, r in by_name.items() if n and not n.isascii())
    cell = ws.cell(rtl, col(ws, 'Business'))
    assert cell.alignment.horizontal == 'right' and cell.alignment.readingOrder == 2
    dq = wb['Disqualified']
    reasons = [dq.cell(r, col(dq, 'Why it was cut')).value for r in range(2, dq.max_row + 1)]
    assert any(x.startswith('chain:') for x in reasons)


def test_read_me(built):
    _p, rows, wb = built
    text = '\n'.join(str(c.value) for row in wb['Read me'].iter_rows() for c in row if c.value)
    listings = sum(len(r['sources']) for r in rows)
    assert f'{listings} listings from 3 sources → {len(rows)} businesses' in text
    assert 'CAN-SPAM' in text and 'HOW THE SCORE WORKS' in text


def test_working_columns_survive_reexport(built):
    project, rows, wb = built
    ws = wb['Call list']
    ws.cell(2, col(ws, 'Notes')).value = 'Call back Tuesday'
    wb.save(project / 'out' / 'leads.xlsx')
    assert export(project).returncode == 0
    ws2 = load_workbook(project / 'out' / 'leads.xlsx')['Call list']
    assert ws2.cell(2, col(ws2, 'Notes')).value == 'Call back Tuesday'


def test_working_columns_follow_a_changed_id(built):
    """A note typed under an id that merged into another business (former_ids) stays with that business."""
    project, rows, wb = built
    ws = wb['Call list']
    old = ws.cell(2, col(ws, 'ID')).value
    ws.cell(2, col(ws, 'Notes')).value = 'Owner prefers WhatsApp'
    wb.save(project / 'out' / 'leads.xlsx')
    doc = json.loads((project / 'data' / 'scored.json').read_text())
    row = next(r for r in doc['items'] if r['id'] == old)
    row['id'], row['former_ids'] = 'b_0000000000ab', [old]
    (project / 'data' / 'scored.json').write_text(json.dumps(doc))
    assert export(project).returncode == 0
    ws2 = load_workbook(project / 'out' / 'leads.xlsx')['Call list']
    r2 = next(i for i in range(2, ws2.max_row + 1) if ws2.cell(i, col(ws2, 'ID')).value == 'b_0000000000ab')
    assert ws2.cell(r2, col(ws2, 'Notes')).value == 'Owner prefers WhatsApp'


def test_missing_sites_and_google(tmp_path):
    def strip(items):
        for r in items:
            r['google'] = None
            r['site_state'] = 'not_crawled' if r.get('website') else 'none'
            r['seo_score'], r['ads'] = None, None
    project, _rows = make_project(tmp_path, strip)
    assert export(project).returncode == 0
    ws = load_workbook(project / 'out' / 'leads.xlsx')['All leads']
    assert {ws.cell(r, col(ws, 'Runs ads')).value for r in range(2, ws.max_row + 1)} == {'not checked'}


@pytest.mark.parametrize('kind', ['phone', 'website', 'instagram', 'place'])
def test_duplicate_gate_refuses(tmp_path, kind):
    def dup(items):
        a, b = items[0], items[1]
        if kind == 'phone':
            b['phones'] = copy.deepcopy(a['phones'])
        elif kind == 'website':
            b['website'], b['website_kind'] = 'https://northwind.example/about', 'site'
        elif kind == 'instagram':
            b['socials']['instagram'] = a['socials']['instagram'].upper() + '/'
        else:
            b['google'] = {**(b['google'] or {}), 'place_id': a['google']['place_id']}
    project, _rows = make_project(tmp_path, dup)
    r = export(project)
    assert r.returncode == 1
    assert 'DUPLICATE CHECK FAILED' in r.stderr and 'Northwind Plumbing' in r.stderr
    assert not (project / 'out' / 'leads.xlsx').exists()


def test_switchboard_phone_allowed(tmp_path):
    def share(items):
        items[1]['phones'] = copy.deepcopy(items[0]['phones'])
        for r in items[:2]:
            r['flags'] = ['switchboard']
    project, _rows = make_project(tmp_path, share)
    assert export(project).returncode == 0
