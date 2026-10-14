#!/usr/bin/env python3
"""
l2p export: data/scored.json -> out/leads.xlsx

    python export_xlsx.py <projectDir>

Four sheets, because the list is worked in three ways and explained once:
  Call list      tiers A and B, not disqualified, best first: what you call this week
  All leads      every live row with the full evidence trail
  Disqualified   what was cut and why, so a judgement call can be reversed
  Read me        what each column means, how the score works, where the rows came from

Before saving, a duplicate gate refuses to write (exit 1) if two rows share a phone number,
an own-site domain, an Instagram handle or a Google place id: the merge stage must fix that.
Working columns (Called on, Outcome, ...) typed into an earlier leads.xlsx are carried over by ID.
No formulas: every value is a measurement taken on the collection date.
"""
import datetime
import hashlib
import json
import os
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PLUGIN = Path(__file__).resolve().parent.parent

INK, LINK = '1A1714', '1D4E89'
BAND = PatternFill('solid', fgColor='F4F3EF')
BOX = Border(bottom=Side(style='thin', color='D3D7D4'))
TIER_FILL = {'A': 'C8E6C9', 'B': 'DCEDC8', 'C': 'FFF9C4', 'D': 'EEEEEE', 'DQ': 'FFCDD2'}
CHECK_FILL = {'MATCH': 'C8E6C9', 'MISMATCH': 'FFCDD2', 'NO PHONE': 'FFE0B2', 'SITE UNREACHABLE': 'FFE0B2'}
RTL = re.compile(r'[֐-ࣿיִ-﷿ﹰ-﻿]')
LATIN = re.compile(r'[A-Za-z]')
SOCIALS = ['instagram', 'facebook', 'linkedin', 'tiktok', 'youtube', 'x', 'whatsapp', 'telegram']
SOCIAL_LABEL = {'x': 'X', 'tiktok': 'TikTok', 'youtube': 'YouTube', 'linkedin': 'LinkedIn', 'whatsapp': 'WhatsApp'}
WORKING = [('Called on', 'w_called_on', 12), ('Outcome', 'w_outcome', 16), ('Next step', 'w_next_step', 26),
           ('Next step date', 'w_next_date', 13), ('Objection', 'w_objection', 26), ('Notes', 'w_notes', 36)]
WEBSITE_TYPE = {'site': 'own website', 'free_subdomain': 'free subdomain (Wix, Weebly...)', 'corporate': 'profile page on a company site',
                'social': 'social page only', 'aggregator': 'platform listing only', 'directory': 'directory page only', 'none': 'no website'}


def font(size=10, bold=False, color=INK):
    return Font(name='Arial', size=size, bold=bold, color=color)


# ---------- config: project.json merged over config/defaults.json ----------
def deep_merge(base, over):
    if not isinstance(base, dict) or not isinstance(over, dict):
        return base if over is None else over
    out = dict(base)
    for k, v in over.items():
        out[k] = deep_merge(base.get(k), v) if isinstance(v, dict) and isinstance(base.get(k), dict) else v
    return out


def load_config(project):
    defaults = json.loads((PLUGIN / 'config' / 'defaults.json').read_text())
    return deep_merge(defaults, json.loads((project / 'project.json').read_text()))


# ---------- helpers ----------
def host(u):
    try:
        p = urlparse(u if re.match(r'^https?://', str(u or ''), re.I) else f'https://{u}')
        h = (p.hostname or '') + (f':{p.port}' if p.port else '')   # keep a non-default port
    except ValueError:
        return ''
    return re.sub(r'^www\d?\.', '', h.lower())


def insta_handle(u):
    u = str(u or '').strip()
    if not u:
        return ''
    if not re.match(r'^https?://', u, re.I):
        return u.lstrip('@').strip('/').lower()
    parts = [p for p in urlparse(u).path.split('/') if p]
    return parts[0].lstrip('@').lower() if parts else ''


def g(r):
    return r.get('google') or {}


def fmt_reasons(lst):
    return '; '.join(f"{x.get('id')} {x.get('points'):+d}" for x in (lst or []) if isinstance(x.get('points'), int))


def value(r, key, work):
    """Display value for one cell. Keys starting with _ are computed."""
    site_state = r.get('site_state') or ('none' if not r.get('website') else 'not_crawled')
    if key.startswith('w_'):
        return work.get(r.get('id'), {}).get(key, '')
    if key == '_sources':
        return ', '.join(dict.fromkeys(s.get('source', '') for s in r.get('sources') or [] if s.get('source')))
    if key == '_source_links':
        return '\n'.join(s.get('url', '') for s in r.get('sources') or [] if s.get('url'))
    if key == '_owner':
        o = r.get('owner') or {}
        if not o.get('name'):
            return 'Not published'
        return o['name'] + (f", {o['role']}" if o.get('role') else '') + (f" ({o['source']})" if o.get('source') else '')
    if key in ('_phone_check', '_phone_site', '_phone_google'):
        pc = r.get('phone_check') or {}
        if key == '_phone_site':
            return pc.get('site', '')
        if key == '_phone_google':
            return pc.get('google', '') or ('not checked' if not r.get('google') else '')
        return (pc.get('verdict') or 'NOT CHECKED') + (f" ({pc['note']})" if pc.get('note') else '')
    if key == '_website_type':
        kind = r.get('website_kind') or ('site' if r.get('website') else 'none')
        return WEBSITE_TYPE.get(kind, kind)
    if key == '_seo_score':
        if r.get('seo_score') is not None:
            return r['seo_score']
        return 'no website' if site_state == 'none' else 'not checked'
    if key == '_seo_band':
        return r.get('seo_band') or ''
    if key == '_seo_issues':
        return '; '.join((r.get('seo_issues') or [])[:3])
    if key.startswith('_social_'):
        return (r.get('socials') or {}).get(key[8:], '') or ''
    if key in ('_g_rating', '_g_reviews'):
        if not r.get('google'):
            return 'not checked'
        v = g(r).get('rating' if key == '_g_rating' else 'reviews')
        return '' if v is None else v
    if key == '_g_maps':
        return g(r).get('maps_url', '') or ''
    if key == '_g_match':
        m = g(r).get('match') or {}
        return f"{m.get('confidence', '')} ({m.get('basis', '')})" if m else ('not checked' if not r.get('google') else 'from Google Maps')
    if key == '_place_id':
        return g(r).get('place_id', '') or ''
    if key == '_ads':
        ads = r.get('ads')
        return 'not checked' if ads is None else (', '.join(ads) if ads else 'No')
    if key == '_defects':
        return '; '.join(x.get('text', '') for x in r.get('need_reasons') or [] if x.get('id') != 'loved_but_weak_site')
    if key in ('_need_reasons', '_interest_reasons', '_negatives'):
        return fmt_reasons(r.get(key[1:]))
    if key == '_dq':
        d = r.get('disqualified') or {}
        return f"{d.get('reason', '')}: {d.get('detail', '')}".strip(': ')
    if key == '_flags':
        return ', '.join(r.get('flags') or [])
    if key == '_email_ok':
        return 'Yes' if r.get('email_ok') else 'No'
    if key == 'email':
        return r.get('email') or '; '.join(r.get('emails') or [])
    v = r.get(key, '')
    if isinstance(v, bool):
        return 'Yes' if v else 'No'
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    return '' if v is None else v


# ---------- columns: (label, key, width) ----------
WORK_COLS = [
    ('Tier', 'tier', 6), ('Priority', 'priority', 9), ('Need', 'need', 7), ('Interest', 'interest', 9),
    ('Business', 'name', 34), ('Category', 'category', 20), ('City', 'city', 14), ('Sources', '_sources', 20),
    ('Owner', '_owner', 32), ('Phone', 'phone', 17), ('Phone check', '_phone_check', 30),
    ('Phone on website', '_phone_site', 22), ('Phone on Google', '_phone_google', 17),
    ('Email', 'email', 30), ('Email found at', 'email_page', 34), ('Website', 'website', 34), ('Website type', '_website_type', 20),
    ('SEO score', '_seo_score', 10), ('SEO band', '_seo_band', 18), ('SEO top issues', '_seo_issues', 50),
] + [(SOCIAL_LABEL.get(s, s.capitalize()), f'_social_{s}', 26) for s in SOCIALS] + [
    ('Google rating', '_g_rating', 9), ('Google reviews', '_g_reviews', 9), ('Google Maps', '_g_maps', 26),
    ('Runs ads', '_ads', 14), ('Incumbent agency', 'incumbent_agency', 22), ('Defects', '_defects', 46),
    ('Why they need you', 'why_they_need_you', 50), ('Why they will care', 'why_they_care', 46), ('Risk', 'risk', 34),
    ('Channel', 'channel', 10), ('Outreach rule', 'outreach_rule', 56),
] + WORKING + [('ID', 'id', 18)]

EXTRA_COLS = [
    ('Address', 'address', 30), ('Region', 'region', 10), ('Postcode', 'postcode', 10), ('Country', 'country', 8),
    ('Category group', 'category_group', 14), ('Ticket', 'ticket', 9), ('Site state', 'site_state', 12), ('Builder', 'builder', 13),
    ('Footer year', 'footer_year', 10), ('SEO measured %', 'seo_measured_pct', 10), ('Email OK', '_email_ok', 9),
    ('Need reasons', '_need_reasons', 44), ('Interest reasons', '_interest_reasons', 44), ('Negatives', '_negatives', 28),
    ('Google match', '_g_match', 18), ('Place ID', '_place_id', 28), ('Source links', '_source_links', 44), ('Flags', '_flags', 16),
    # Same key as the call list: dropped by uniq() rather than printed twice.
    ('Incumbent agency', 'incumbent_agency', 22),
]

DQ_COLS = [('Business', 'name', 34), ('Category', 'category', 20), ('City', 'city', 14), ('Why it was cut', '_dq', 50),
           ('Priority', 'priority', 9), ('Phone', 'phone', 17), ('Website', 'website', 34), ('Sources', '_sources', 20), ('ID', 'id', 18)]

HYPER = {'website', '_g_maps', 'email_page'} | {f'_social_{s}' for s in SOCIALS}
WRAP = {'name', '_owner', '_phone_check', '_seo_issues', '_defects', 'why_they_need_you', 'why_they_care', 'risk',
        'outreach_rule', '_need_reasons', '_interest_reasons', '_negatives', '_source_links', '_dq', 'address', 'w_notes'}
NUM = {'priority', 'need', 'interest', '_seo_score', '_g_rating', '_g_reviews', 'footer_year', 'seo_measured_pct'}


def uniq(cols):
    """Drop a column whose key already appeared (the old duplicated-column bug)."""
    out, seen = [], set()
    for c in cols:
        if c[1] not in seen:
            out.append(c)
            seen.add(c[1])
    labels = [c[0] for c in out]
    assert len(set(labels)) == len(labels), f'duplicate column labels: {labels}'
    return out


# ---------- duplicate gate ----------
def duplicates(rows):
    keys = {}
    for r in rows:
        switchboard = 'switchboard' in (r.get('flags') or [])
        ks = [(f"phone {p['key']}", switchboard) for p in r.get('phones') or [] if p.get('key')]
        kind = r.get('website_kind') or ('site' if r.get('website') else 'none')
        if r.get('website') and kind in ('site', 'free_subdomain') and host(r['website']):
            ks.append((f"website {host(r['website'])}", False))
        h = insta_handle((r.get('socials') or {}).get('instagram'))
        if h:
            ks.append((f'instagram {h}', False))
        if g(r).get('place_id'):
            ks.append((f"google place {g(r)['place_id']}", False))
        for k, shared_ok in dict(ks).items():
            keys.setdefault(k, []).append((r, shared_ok))
    # A switchboard number shared only by rows flagged as such is allowed (flagged in merge).
    return [(k, [r for r, _ in v]) for k, v in keys.items()
            if len({id(r) for r, _ in v}) > 1 and not all(ok for _, ok in v)]


# ---------- carry working columns over from the previous workbook ----------
def id_aliases(r):
    """Ids this business may have had (lib/dedupe.mjs idAliases): ids merged into it, and the id each of its
    listings would have had alone (1.0.0 re-derived ids from the first source)."""
    out = list(r.get('former_ids') or [])
    out += ['b_' + hashlib.sha1(str(s.get('id')).encode('utf-8')).hexdigest()[:12] for s in r.get('sources') or [] if s.get('id')]
    return [a for a in dict.fromkeys(out) if a and a != r.get('id')]


def carry_work(rows, work):
    """Notes typed under an id the business no longer has follow it to its current id."""
    current = {r.get('id') for r in rows}
    for r in rows:
        if r.get('id') in work:
            continue
        old = next((a for a in id_aliases(r) if a not in current and a in work), None)
        if old:
            work[r['id']] = work[old]
    return work


def previous_work(path):
    work = {}
    if not path.exists():
        return work
    try:
        wb = load_workbook(path, read_only=True)
    except Exception:                       # noqa: BLE001 - an unreadable old file must not block a new one
        return work
    label_to_key = {label: key for label, key, _w in WORKING}
    for name in ('All leads', 'Call list'):        # call list last: that is where people type
        if name not in wb.sheetnames:
            continue
        it = wb[name].iter_rows(values_only=True)
        head = next(it, None) or []
        if 'ID' not in head:
            continue
        idx = head.index('ID')
        cols = [(i, label_to_key[h]) for i, h in enumerate(head) if h in label_to_key]
        for row in it:
            if not row or idx >= len(row) or not row[idx]:
                continue
            for i, k in cols:
                if i < len(row) and row[i] not in (None, ''):
                    work.setdefault(row[idx], {})[k] = row[i]
    wb.close()
    return work


# ---------- sheets ----------
def safe(v):
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub('', v)
        if v.startswith('='):
            v = "'" + v                     # never let a scraped value become a formula
    return v


def is_rtl(v):
    if not isinstance(v, str):
        return False
    n = len(RTL.findall(v))
    return n > 0 and n * 2 >= len(LATIN.findall(v))


def header_style(color):
    hexv = re.sub(r'[^0-9A-Fa-f]', '', str(color or ''))[:6]
    hexv = hexv if len(hexv) == 6 else '1F4E79'
    rgb = [int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    lum = 0.2126 * rgb[0] + 0.7152 * rgb[1] + 0.0722 * rgb[2]
    return PatternFill('solid', fgColor=hexv.upper()), ('FFFFFF' if lum < 0.6 else INK)


def sheet(wb, title, cols, data, head, work):
    fill, ink = head
    ws = wb.create_sheet(title)
    for i, (label, _k, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=label)
        c.font, c.fill = font(10, True, ink), fill
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for ri, r in enumerate(data, 2):
        for ci, (_l, key, _w) in enumerate(cols, 1):
            v = safe(value(r, key, work))
            c = ws.cell(row=ri, column=ci, value=v)
            c.font, c.border = font(), BOX
            rtl = is_rtl(v)
            c.alignment = Alignment(vertical='top', wrap_text=key in WRAP or rtl,
                                    horizontal='right' if rtl or (key in NUM and isinstance(v, (int, float))) else None,
                                    readingOrder=2 if rtl else 0)
            if ri % 2 == 0:
                c.fill = BAND
            if key == 'tier':
                t = 'DQ' if r.get('disqualified') else r.get('tier')
                c.fill = PatternFill('solid', fgColor=TIER_FILL.get(t, 'EEEEEE'))
                c.font, c.alignment = font(10, True), Alignment(horizontal='center', vertical='top')
            elif key == 'priority':
                c.font = font(10, True)
            elif key == '_phone_check' and v:
                for prefix, color in CHECK_FILL.items():
                    if str(v).startswith(prefix):
                        c.fill = PatternFill('solid', fgColor=color)
            if key in HYPER and isinstance(v, str) and re.match(r'^https?://', v) and len(v) < 2000:
                c.hyperlink, c.font = v, font(color=LINK)
            elif key == 'email' and isinstance(v, str) and re.match(r'^[^@\s;]+@[^@\s;]+\.[a-z]{2,}', v, re.I):
                c.hyperlink, c.font = 'mailto:' + re.split(r'[;\s]', v)[0], font(color=LINK)
    ws.freeze_panes = 'F2' if len(cols) > 5 else 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{max(2, len(data) + 1)}'
    return ws


COMPLIANCE = {
    'CASL': 'CASL (Canada): email only addresses the business published on its own website, with no notice refusing '
            'unsolicited email ("Email OK" column). Directory emails do not count as consent. Business phone numbers are '
            'exempt from the National DNCL, but keep your own do-not-call list. Robocalls need prior consent.',
    'GDPR': 'GDPR (EU) plus national law: call or write first for cold B2B outreach; email only after they agree. Germany and '
            'Austria (UWG): no cold email or fax, and call only where interest can be presumed; a letter is the safe first touch. '
            'Owner names and sole traders\' numbers are personal data: document your legitimate interest and tell people where '
            'you got their details.',
    'PECR': 'PECR + UK-GDPR: email to companies (Ltd, LLP, plc) is allowed with your identity and an opt-out; sole traders and '
            'partnerships need prior consent. Check TPS/CTPS before calling.',
    'CAN-SPAM': 'CAN-SPAM (US): business email is allowed if every message carries your real identity and postal address, an '
                'honest subject line, and an opt-out you honour within 10 business days. Respect sites that refuse unsolicited email.',
    'SPAM_ACT': 'Spam Act (Australia): email only addresses conspicuously published by the business, relevant to its role, with '
                'no notice refusing unsolicited email. Identify yourself and include an unsubscribe.',
}


def regime_of(s):
    x = str(s or '').upper()
    for k in ('CASL', 'PECR', 'GDPR'):
        if k in x:
            return k
    if 'SPAM ACT' in x or 'SPAM-ACT' in x:
        return 'SPAM_ACT'
    return 'CAN-SPAM' if re.sub(r'[^A-Z]', '', x) == 'CANSPAM' else 'UNKNOWN'


def readme(wb, cfg, rows, updated):
    ws = wb.create_sheet('Read me')
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 110
    live = [r for r in rows if not r.get('disqualified')]
    t = lambda k: sum(1 for r in live if r.get('tier') == k)       # noqa: E731
    per_source = {}
    for r in rows:
        for s in r.get('sources') or []:
            per_source[s.get('source', '?')] = per_source.get(s.get('source', '?'), 0) + 1
    listings = sum(per_source.values())
    google = sum(1 for r in rows if r.get('google'))
    crawled = sum(1 for r in rows if r.get('site_state') not in (None, '', 'none', 'not_crawled'))
    seo = sum(1 for r in rows if r.get('seo_score') is not None)
    own = sum(1 for r in rows if r.get('website_kind') in ('site', 'free_subdomain', 'corporate'))
    pc = {}
    for r in rows:
        k = (r.get('phone_check') or {}).get('verdict') or 'NOT CHECKED'
        pc[k] = pc.get(k, 0) + 1
    dq = {}
    for r in rows:
        if r.get('disqualified'):
            k = r['disqualified'].get('reason', '?')
            dq[k] = dq.get(k, 0) + 1
    regime = cfg.get('compliance', {}).get('regime', '')
    tiers = {'A': 55, 'B': 38, 'C': 24, **(cfg.get('scoring', {}).get('tiers') or {})}
    agency = cfg.get('agency', {}).get('name') or ''
    lines = [
        (f"{cfg.get('name') or 'Leads'}: leads" + (f' for {agency}' if agency else ''), ''), ('', ''),
        ('Date', datetime.date.today().isoformat() + (f' (scored {updated})' if updated else '')),
        ('Businesses', f'{listings:,} listings from {len(per_source)} source{"s" if len(per_source) != 1 else ""} '
                       f'→ {len(rows):,} businesses'),
        ('Rows per source', ' · '.join(f'{k} {v:,}' for k, v in sorted(per_source.items(), key=lambda x: -x[1])) or '-'),
        ('Worth working (A + B)', f"{t('A') + t('B'):,}   (A {t('A')}, B {t('B')}, C {t('C')}, D {t('D')}, "
                                  f"disqualified {len(rows) - len(live)})"),
        ('Disqualified, by reason', ' · '.join(f'{k} {v}' for k, v in sorted(dq.items(), key=lambda x: -x[1])) or 'none'),
        ('Websites', f'{own:,} have a website of their own, {crawled:,} crawled, {seo:,} with an SEO score.'),
        ('Google', f'{google:,} of {len(rows):,} matched to a Google profile. "not checked" means no Google data for that row, '
                   'not that the business has no Google profile.'),
        ('Phone check results', ' · '.join(f'{k} {v:,}' for k, v in sorted(pc.items(), key=lambda x: -x[1]))),
        ('', ''),
        ('HOW THE SCORE WORKS', ''),
        ('Need', 'What a stranger can see is wrong with their web presence, 0-100: no website or only a social page, a site '
                 'that does not load, not mobile-friendly, no secure connection, no booking or quote form, a phone you cannot tap, '
                 'an old footer, a dead analytics tag, weak SEO, a website phone that differs from Google, and "customers love '
                 'them but the website does not show it".'),
        ('Interest', 'Whether they will act, 0-100: ad tags on the site (they already pay for customers), paid chat or booking '
                     'software, replies to reviews, new reviews every month, an Instagram profile, a failed DIY site, high-value '
                     'customers, a named owner, listed on several sources. An agency credited on their site, a dormant profile, a '
                     '"no unsolicited email" notice or several locations on one site count against.'),
        ('Priority', 'Square root of Need x Interest, 0-100. A geometric mean, so a lead strong on one axis and empty on the other '
                     'ranks below a balanced one.'),
        ('Tier', f"A ≥ {tiers['A']}, B ≥ {tiers['B']}, C ≥ {tiers['C']}, D below. A and B are on the Call list; C and D are in "
                 'All leads. Disqualified rows (closed, chains, competitors, unreachable, already well served) are on their own '
                 'sheet with the reason.'),
        ('Signals', 'Every point is listed in All leads > Need reasons / Interest reasons / Negatives. A signal only counts when '
                    'its data exists: no crawl, no website points; no Google match, no Google points. Weights and the full list: '
                    'docs/SCORING.md in the lead-to-proposal repo; change them in project.json > scoring.weights.'),
        ('', ''),
        ('COLUMNS', ''),
        ('Sources', 'Every source the business was found on. All leads > Source links has each listing.'),
        ('Owner', 'Name, role and where it was found. "Not published" means none was found.'),
        ('Phone check', 'Compares the numbers on their own website with their Google profile. MATCH: they agree. MISMATCH: both are '
                        'shown so you can see which is out of date. NO PHONE ON WEBSITE: the site shows none. SITE ONLY: Google lists '
                        'no number. SITE UNREACHABLE: the site did not load. NOT CHECKED: not crawled yet, crawl blocked, or no Google '
                        'match (the note says which).'),
        ('Email / Email found at', 'The best address and the page it was published on.'),
        ('SEO score', '0-100 from a basic on-page audit of the home page and a few internal pages. "not checked" = not crawled.'),
        ('Runs ads', 'Google Ads, Meta Pixel or TikTok tags found on their site. "not checked" = the site was not reached.'),
        ('Defects', 'Every observed problem, strongest first.'),
        ('Why they need you / will care / Risk', 'Plain sentences built from the strongest reasons: say these on the call.'),
        ('Channel / Outreach rule', 'How to make first contact under your compliance regime, and the rule that applies.'),
        ('Called on ... Notes', 'Yours to fill in. Re-running `l2p export` keeps what you typed (matched by ID).'),
        ('', ''),
        ('BEFORE YOU CONTACT ANYONE', ''),
        ('Compliance', COMPLIANCE.get(regime_of(regime), f'Regime "{regime}" is not one the tool knows: call or visit first and '
                                                         'check your local law before emailing.')),
        ('Nothing is sent', 'This workbook is a list. Every call, email and visit is made by you.'),
        ('Not legal advice', 'Check the rules that apply to you before contacting anyone.'),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ca, cb = ws.cell(row=i, column=1, value=a), ws.cell(row=i, column=2, value=b)
        ca.font = font(14, True) if i == 1 else font(11, True) if (a and a.isupper()) else font(10, True)
        ca.alignment = Alignment(vertical='top')
        cb.font, cb.alignment = font(), Alignment(wrap_text=True, vertical='top')
    ws.sheet_view.showGridLines = False


def main(argv):
    if not argv:
        print('Usage: export_xlsx.py <projectDir>', file=sys.stderr)
        return 2
    project = Path(argv[0]).resolve()
    scored_path = project / 'data' / 'scored.json'
    if not scored_path.exists():
        print('No data/scored.json yet. Run `l2p score` first.', file=sys.stderr)
        return 2
    cfg = load_config(project)
    doc = json.loads(scored_path.read_text())
    rows = doc.get('items', doc) if isinstance(doc, dict) else doc

    dupes = duplicates(rows)
    if dupes:
        print(f'DUPLICATE CHECK FAILED: {len(dupes)} key(s) appear on more than one row. Nothing was saved.\n'
              'Fix the merge (l2p merge, see data/merge-report.json), re-score, then export again.', file=sys.stderr)
        for k, rs in dupes[:50]:
            print(f'  {k}: ' + '  <->  '.join(f'"{r.get("name")}" ({r.get("id")})' for r in rs), file=sys.stderr)
        return 1

    out = project / 'out' / 'leads.xlsx'
    work = carry_work(rows, previous_work(out))
    head = header_style(cfg.get('agency', {}).get('color'))
    rank = lambda r: (-(r.get('priority') or 0), -(r.get('need') or 0), str(r.get('name') or ''))   # noqa: E731
    live = sorted([r for r in rows if not r.get('disqualified')], key=rank)
    call = [r for r in live if r.get('tier') in ('A', 'B')]
    dq = sorted([r for r in rows if r.get('disqualified')], key=rank)

    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'Call list', uniq(WORK_COLS), call, head, work)
    sheet(wb, 'All leads', uniq(WORK_COLS + EXTRA_COLS), live, head, work)
    sheet(wb, 'Disqualified', uniq(DQ_COLS), dq, head, work)
    readme(wb, cfg, rows, doc.get('updated', '') if isinstance(doc, dict) else '')
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(f'.{out.name}.{os.getpid()}.tmp')
    wb.save(tmp)
    os.replace(tmp, out)
    print(f'Duplicate check passed ({len(rows):,} rows).')
    print(f'{out}: call list {len(call):,}, all leads {len(live):,}, disqualified {len(dq):,}')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
